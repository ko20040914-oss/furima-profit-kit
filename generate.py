"""フリマ・物販向けの数式付き Excel キットを生成する。"""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

SHEETS = ['はじめに', '販売記録', '手数料マスタ', '在庫一覧', '採算シミュレータ', '月次集計']
MONEY = '#,##0"円";[Red]-#,##0"円";0"円"'
NAVY = '17324D'
TEAL = '087F8C'
INPUT = 'EDF6FF'
OUTPUT = 'EAF5EF'


def _base(ws, title, subtitle, headers, widths):
    ws.sheet_view.showGridLines = False
    ws['A1'] = title
    ws['A1'].font = Font(name='Yu Gothic', size=20, bold=True, color=NAVY)
    ws['A2'] = subtitle
    ws['A2'].font = Font(name='Yu Gothic', size=10, color='526577')
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 26
    for i, text in enumerate(headers, 1):
        cell = ws.cell(4, i, text)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Yu Gothic', bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 30
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'C5' if ws.title == '販売記録' else 'A5'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'
    ws.sheet_properties.tabColor = TEAL


def _cell_style(cell, formula=False, number_format=None):
    cell.font = Font(name='Yu Gothic', size=10, color='17324D' if formula else '145DA0')
    cell.fill = PatternFill('solid', fgColor=OUTPUT if formula else INPUT)
    cell.alignment = Alignment(vertical='center')
    if number_format:
        cell.number_format = number_format


def _validation(ws, kind, formula1, ranges, formula2=None):
    dv = DataValidation(type=kind, formula1=formula1, formula2=formula2, allow_blank=True)
    dv.errorTitle = '入力値をご確認ください'
    dv.error = '候補または指定範囲の値を入力してください。'
    dv.showErrorMessage = True
    dv.errorStyle = 'stop'
    ws.add_data_validation(dv)
    for target in ranges:
        dv.add(target)


def build_workbook(rows=1000, year=2026):
    if not 3 <= rows <= 50000:
        raise ValueError('rows は 3～50000 の範囲で指定してください。')
    if not 1900 <= year <= 9998:
        raise ValueError('year は 1900～9998 の範囲で指定してください。')
    end = rows + 4
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        wb.create_sheet(name)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    wb.properties.title = '実利益・在庫管理キット'
    wb.properties.subject = 'フリマ・物販の実利益と在庫管理'
    intro = wb['はじめに']
    _base(intro, '実利益・在庫管理キット', '販売の手残りと在庫の滞留を、ひとつのファイルで。', ['手順・項目', '使い方'], [28, 108])
    instructions = [
        ('1 ｜ 手数料を設定', '「手数料マスタ」で手数料率を確認・編集します。初期値は例です。契約・キャンペーン等の実際の料率を確認してください。'),
        ('2 ｜ 仕入れを記録', '「販売記録」に商品名・プラットフォーム・仕入価格・仕入日を入力し、状態を「在庫」にします。'),
        ('3 ｜ 売れたら更新', '販売日・販売価格・送料・梱包費を入力し、状態を「販売済み」にします。費用がない場合も 0 を入力します。'),
        ('4 ｜ 利益を確認', '販売手数料＝販売価格×マスタ料率（円未満四捨五入）。利益＝販売価格－販売手数料－送料－仕入価格－梱包費。'),
        ('5 ｜ 在庫を確認', '「在庫一覧」は状態が在庫の商品を自動表示します。経過日数は今日－仕入日。30日以上を色付けします。'),
        ('6 ｜ 採算を検討', '「採算シミュレータ」の青いセルを変更します。手残りは仕入・送料・梱包費も差し引いた利益です。利益率＝利益÷売値。'),
        ('7 ｜ 月次を確認', '「月次集計」は販売日の属する月で、販売済みの商品のみ集計します。対象年は月次集計 B2 で変更できます。'),
        ('入力と計算の色', '青色セルは入力欄、薄緑色セルは数式です。数式は上書きせず、入力欄のみ編集してください。'),
        ('サンプル', '商品A・商品B（販売済み）と商品C（在庫）を収録。使い始める前に入力欄を消去・置換してください。'),
        ('継続利用・行数', f'販売記録・在庫一覧は各 {rows:,} 件分の数式を用意。空いている入力行に追記できます。'),
        ('上限を超えるとき', '既存データをバックアップし、付属の extend.py で入力済みファイルを別名に拡張できます。数式・集計範囲・入力規則も更新します。'),
        ('新しい大容量テンプレート', 'Python の --rows で容量を増やした新規テンプレートを作成できます。再生成は既存データを引き継がないため別名保存してください。'),
        ('料率変更の注意', 'マスタ変更は過去の販売記録も再計算します。料率が変わるときは「メルカリ旧料率」等の別名を追加し、過去の料率を残してください。'),
        ('未登録・入力不足', '未登録・重複プラットフォームや未入力・不正な料率は「要確認」と表示。必須の費用等が空欄だと利益は空欄になります。月次は入力完了後に確認してください。'),
        ('対象外', '税金・振込手数料・返品・値引き・複数商品の送料按分等は含みません。必要な費用は各入力額に反映してください。'),
        ('Excel / Google', 'Excelで開くか、GoogleドライブへアップロードしてGoogleスプレッドシートで開いてください。マクロ・外部リンクはありません。'),
    ]
    for r, (label, body) in enumerate(instructions, 5):
        intro.cell(r, 1, label)
        intro.cell(r, 2, body)
        for c in intro[r][:2]:
            _cell_style(c, True)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        intro.row_dimensions[r].height = 42
    master = wb['手数料マスタ']
    _base(master, '手数料マスタ', '編集可能な参考値です。実際の契約・最新の料率を確認してください。変更は過去分にも反映されます。', ['プラットフォーム', '手数料率', '備考'], [30, 18, 72])
    for r in range(5, 105):
        for c in range(1, 4):
            _cell_style(master.cell(r, c), number_format='0.0%' if c == 2 else None)
    for r, data in enumerate([('メルカリ', .1, '参考初期値。販売価格に対する割合。'), ('Yahoo!フリマ', .05, '参考例。キャンペーン等による変更を確認。'), ('その他', .1, '仮の料率。利用サービスに合わせて編集。')], 5):
        for c, value in enumerate(data, 1):
            master.cell(r, c, value)
    named = DefinedName('PlatformList', attr_text="'手数料マスタ'!$A$5:$A$104")
    if hasattr(wb.defined_names, 'add'):
        wb.defined_names.add(named)
    else:
        wb.defined_names.append(named)
    _validation(master, 'decimal', '0', ['B5:B104'], '1')
    master.auto_filter.ref = 'A4:C104'
    sales = wb['販売記録']
    _base(sales, '販売記録', f'青＝入力 ｜ 緑＝自動計算 ｜ {rows:,}件分 ｜ 費用ゼロは 0 を入力 ｜ 販売日と仕入日は別に記録', ['販売日', '商品名', 'プラットフォーム', '仕入価格', '販売価格', '販売手数料', '送料', '梱包費', '利益', '状態', '仕入日', '在庫連番'], [15, 26, 24, 16, 16, 18, 15, 15, 17, 16, 15, 12])
    samples = [
        [date(2026, 10, 1), '商品A', 'メルカリ', 2000, 5000, None, 700, 100, None, '販売済み', date(2026, 9, 15)],
        [date(2026, 10, 2), '商品B', 'メルカリ', 500, 1500, None, 210, 50, None, '販売済み', date(2026, 9, 20)],
        [None, '商品C', 'メルカリ', 1000, None, None, None, None, None, '在庫', date(2026, 9, 1)],
    ]
    for r in range(5, end + 1):
        for c in range(1, 13):
            fmt = 'yyyy/mm/dd' if c in (1, 11) else MONEY if c in range(4, 10) else None
            _cell_style(sales.cell(r, c), c in (6, 9, 12), fmt)
        sales.cell(r, 6, f'=IF(OR(B{r}="",J{r}<>"販売済み",E{r}="",C{r}=""),"",IFERROR(IF(OR(COUNTIF(\'手数料マスタ\'!$A$5:$A$104,C{r})<>1,COUNTIFS(\'手数料マスタ\'!$A$5:$A$104,C{r},\'手数料マスタ\'!$B$5:$B$104,"<>")<>1,NOT(ISNUMBER(VLOOKUP(C{r},\'手数料マスタ\'!$A$5:$B$104,2,FALSE))),VLOOKUP(C{r},\'手数料マスタ\'!$A$5:$B$104,2,FALSE)<0,VLOOKUP(C{r},\'手数料マスタ\'!$A$5:$B$104,2,FALSE)>1),"要確認",ROUND(E{r}*VLOOKUP(C{r},\'手数料マスタ\'!$A$5:$B$104,2,FALSE),0)),"要確認"))')
        sales.cell(r, 9, f'=IF(OR(B{r}="",J{r}<>"販売済み",D{r}="",E{r}="",G{r}="",H{r}=""),"",IF(ISNUMBER(F{r}),E{r}-D{r}-F{r}-G{r}-H{r},"要確認"))')
        sales.cell(r, 12, f'=IF(AND(B{r}<>"",J{r}="在庫"),COUNTIFS($J$5:J{r},"在庫",$B$5:B{r},"<>"),"")')
    for r, data in enumerate(samples, 5):
        for c, value in enumerate(data, 1):
            if c not in (6, 9):
                sales.cell(r, c, value)
    sales.column_dimensions['L'].hidden = True
    sales.auto_filter.ref = f'A4:K{end}'
    _validation(sales, 'list', 'PlatformList', [f'C5:C{end}'])
    _validation(sales, 'list', '"販売済み,在庫"', [f'J5:J{end}'])
    _validation(sales, 'decimal', '0', [f'{c}5:{c}{end}' for c in ['D','E','G','H']], '1000000000')
    sales.conditional_formatting.add(f'I5:I{end}', CellIsRule(operator='lessThan', formula=['0'], fill=PatternFill('solid', fgColor='FDE7E9')))
    inv = wb['在庫一覧']
    _base(inv, '在庫一覧', '販売記録の「在庫」を自動表示 ｜ 経過30日以上は要チェック ｜ 販売記録の入力欄で更新します', ['仕入日', '商品名', 'プラットフォーム', '仕入価格', '経過日数', '参照行'], [17, 30, 26, 19, 18, 12])
    for r in range(5, end + 1):
        inv.cell(r, 6, f'=IFERROR(MATCH(ROW()-4,\'販売記録\'!$L$5:$L${end},0),"")')
        for c, source in [(1, 'K'), (2, 'B'), (3, 'C'), (4, 'D')]:
            ref = f"INDEX('販売記録'!${source}$5:${source}${end},$F{r})"
            inv.cell(r, c, f'=IF($F{r}="","",IF({ref}="","",{ref}))')
        inv.cell(r, 5, f'=IF(OR(A{r}="",B{r}=""),"",MAX(0,TODAY()-A{r}))')
        for c in range(1, 7):
            _cell_style(inv.cell(r, c), True, 'yyyy/mm/dd' if c == 1 else MONEY if c == 4 else '0' if c == 5 else None)
    inv.column_dimensions['F'].hidden = True
    inv.auto_filter.ref = f'A4:E{end}'
    inv.conditional_formatting.add(f'E5:E{end}', CellIsRule(operator='greaterThanOrEqual', formula=['30'], fill=PatternFill('solid', fgColor='FFF0C2')))
    sim = wb['採算シミュレータ']
    _base(sim, '採算シミュレータ', '青いセルを変更して採算を確認 ｜ 手残りは全入力費用を引いた利益 ｜ 料率は直接入力', ['項目', '金額・割合', '説明'], [28, 24, 76])
    for r, label, value in [(5,'売値',3000),(6,'仕入',1200),(7,'送料',520),(8,'梱包費',80),(9,'手数料率',.1)]:
        sim.cell(r,1,label)
        sim.cell(r,2,value)
        _cell_style(sim.cell(r,2), False, '0.0%' if r == 9 else MONEY)
    formulas = {
        11: ('販売手数料', '=IF(COUNT(B5:B9)<5,"",ROUND(B5*B9,0))', '売値×手数料率。円未満四捨五入。'),
        12: ('手残り', '=IF(COUNT(B5:B9)<5,"",B5-SUM(B6:B8)-B11)', '売値－仕入－送料－梱包費－販売手数料。'),
        13: ('利益率', '=IF(OR(COUNT(B5:B9)<5,B5=0),"",B12/B5)', '手残り÷売値。売値ゼロの場合は空欄。'),
        14: ('損益分岐売値', '=IF(COUNT(B6:B9)<4,"",IF(B9>=1,"算出不可",ROUNDUP(SUM(B6:B8)/(1-B9),0)))', '手数料の端数丸め前の費用回収目安（円単位切上げ）。料率100%は算出不可。'),
    }
    for r, (label, formula, note) in formulas.items():
        sim.cell(r,1,label)
        sim.cell(r,2,formula)
        sim.cell(r,3,note)
        _cell_style(sim.cell(r,2), True, '0.0%' if r == 13 else MONEY)
    _validation(sim,'decimal','0',['B5:B8'],'1000000000')
    _validation(sim,'decimal','0',['B9'],'1')
    monthly = wb['月次集計']
    _base(monthly, '月次集計', '対象年', ['月', '売上', '手数料', '送料', '仕入', '梱包費', '利益'], [18, 20, 20, 20, 20, 20, 21])
    monthly['B2'] = year
    _cell_style(monthly['B2'], False, '0"年"')
    monthly['C2'] = '販売済みのみ・販売日基準。未入力の費用がないか販売記録も確認してください。'
    _validation(monthly,'whole','1900',['B2'],'9998')
    for r in range(5, 17):
        monthly.cell(r,1,f'=DATE($B$2,{r-4},1)')
        _cell_style(monthly.cell(r,1),True,'yyyy"年"m"月"')
        for col, source in enumerate(['E','F','G','D','H','I'],2):
            monthly.cell(r,col,f'=SUMIFS(\'販売記録\'!${source}$5:${source}${end},\'販売記録\'!$J$5:$J${end},"販売済み",\'販売記録\'!$A$5:$A${end},">="&$A{r},\'販売記録\'!$A$5:$A${end},"<"&DATE(YEAR($A{r}),MONTH($A{r})+1,1))')
            _cell_style(monthly.cell(r,col),True,MONEY)
    monthly['A18'] = '年間合計'
    for c in range(2,8):
        letter = get_column_letter(c)
        monthly.cell(18,c,f'=SUM({letter}5:{letter}16)')
        _cell_style(monthly.cell(18,c),True,MONEY)
        monthly.cell(18,c).font = Font(name='Yu Gothic', bold=True, color=TEAL, size=12)
    for ws in wb:
        ws.sheet_properties.outlinePr.summaryRight = False
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=Path('output/実利益・在庫管理キット.xlsx'))
    parser.add_argument('--rows', type=int, default=1000)
    parser.add_argument('--year', type=int, default=2026)
    args = parser.parse_args()
    wb = build_workbook(args.rows, args.year)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
