"""Integration checks evaluate the generated Excel formulas in LibreOffice.

Run: python -m unittest discover -s tests -v
Requires the soffice executable; tests never replace Excel formulas with Python values.
"""
from datetime import datetime, date
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

from openpyxl import load_workbook
from generate import build_workbook


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if not cls.soffice:
            raise RuntimeError('LibreOffice (soffice) is required for formula verification')

    def recalculate(self, book):
        """Open and export XLSX with a separate user profile to force calculation."""
        with tempfile.TemporaryDirectory(prefix='furima-test-') as tmp:
            root = Path(tmp)
            source = root / 'input.xlsx'
            out = root / 'recalculated'
            out.mkdir()
            book.save(source)
            command = [self.soffice, '-env:UserInstallation=' + (root / 'profile').as_uri(),
                       '--headless', '--convert-to', 'xlsx', '--outdir', str(out), str(source)]
            run = subprocess.run(command, capture_output=True, text=True, timeout=90)
            result = out / source.name
            self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
            self.assertTrue(result.is_file(), run.stdout + run.stderr)
            values = load_workbook(result, data_only=True)
            formulas = load_workbook(result, data_only=False)
            for sheet in values:
                for row in sheet:
                    for cell in row:
                        self.assertNotEqual(cell.data_type, 'e', f'{sheet.title}!{cell.coordinate}: {cell.value}')
            self.assertEqual(formulas['販売記録']['I5'].data_type, 'f')
            return values

    def test_acceptance_examples_and_inventory(self):
        book = build_workbook(rows=20, year=2026)
        self.assertEqual(book.sheetnames, ['はじめに', '販売記録', '手数料マスタ', '在庫一覧', '採算シミュレータ', '月次集計'])
        for coordinate in ['F5', 'I5', 'F24', 'I24']:
            self.assertEqual(book['販売記録'][coordinate].data_type, 'f')
        values = self.recalculate(book)
        sales = values['販売記録']
        self.assertEqual(sales['F5'].value, 500)
        self.assertEqual(sales['I5'].value, 1700)
        self.assertEqual(sales['F6'].value, 150)
        self.assertEqual(sales['I6'].value, 590)
        monthly = values['月次集計']
        self.assertEqual(monthly['B18'].value, 6500)
        self.assertEqual(monthly['G18'].value, 2290)
        self.assertEqual(monthly['B14'].value, 6500)
        simulator = values['採算シミュレータ']
        self.assertEqual(simulator['B11'].value, 300)
        self.assertEqual(simulator['B12'].value, 900)
        self.assertAlmostEqual(simulator['B13'].value, .3)
        self.assertAlmostEqual(simulator['B14'].value, 2000)
        inventory = values['在庫一覧']
        self.assertEqual(inventory['B5'].value, '商品C')
        self.assertEqual(inventory['D5'].value, 1000)
        age = (date.today() - date(2026, 9, 1)).days
        self.assertIn(inventory['E5'].value, [age - 1, age, age + 1])
        self.assertIn(inventory['B6'].value, [None, ''])
        self.assertIn(sales['I7'].value, [None, ''])
        self.assertIn(sales['I24'].value, [None, ''])

    def test_master_edit_and_new_row_recalculate(self):
        book = build_workbook(rows=20, year=2026)
        book['手数料マスタ']['B5'] = .2
        sales = book['販売記録']
        for col, value in {'A': datetime(2026, 11, 1), 'B': '追加商品', 'C': 'メルカリ',
                           'D': 100, 'E': 1000, 'G': 100, 'H': 20, 'J': '販売済み',
                           'K': datetime(2026, 10, 1)}.items():
            sales[f'{col}24'] = value
        values = self.recalculate(book)
        self.assertEqual(values['販売記録']['F5'].value, 1000)
        self.assertEqual(values['販売記録']['I5'].value, 1200)
        self.assertEqual(values['販売記録']['F24'].value, 200)
        self.assertEqual(values['販売記録']['I24'].value, 580)
        self.assertEqual(values['月次集計']['B15'].value, 1000)
        self.assertEqual(values['月次集計']['G18'].value, 2220)

    def test_dates_unsold_and_inventory_status_changes(self):
        book = build_workbook(rows=20, year=2026)
        sales = book['販売記録']
        sales['A5'] = datetime(2025, 12, 31)
        sales['A6'] = datetime(2026, 1, 1)
        # Unsold stock with a proposed selling price must not count as sales.
        sales['A7'] = datetime(2026, 1, 31)
        sales['E7'] = 9000
        # A sold inventory item disappears; another stock item is compacted to row 5.
        sales['J7'] = '販売済み'
        sales['G7'] = 0
        sales['H7'] = 0
        for col, value in {'B': '追加在庫', 'C': 'メルカリ', 'D': 200, 'E': 9999,
                           'J': '在庫', 'K': datetime(2026, 10, 1),
                           'A': datetime(2026, 1, 31)}.items():
            sales[f'{col}8'] = value
        values = self.recalculate(book)
        self.assertEqual(values['月次集計']['B5'].value, 10500)
        self.assertEqual(values['月次集計']['B18'].value, 10500)
        self.assertEqual(values['在庫一覧']['B5'].value, '追加在庫')
        self.assertIn(values['在庫一覧']['B6'].value, [None, ''])
        self.assertIn(values['販売記録']['I8'].value, [None, ''])

    def test_unknown_platform_and_simulator_zero_sales_are_safe(self):
        book = build_workbook(rows=20, year=2026)
        book['販売記録']['C5'] = '未登録サービス'
        book['採算シミュレータ']['B5'] = 0
        values = self.recalculate(book)
        self.assertEqual(values['販売記録']['F5'].value, '要確認')
        self.assertNotIsInstance(values['販売記録']['I5'].value, (int, float))
        self.assertNotIsInstance(values['採算シミュレータ']['B13'].value, (int, float))

    def test_blank_rate_and_duplicate_platform_are_flagged(self):
        book = build_workbook(rows=20, year=2026)
        master = book['手数料マスタ']
        master['B5'] = None
        master['A100'] = '重複サービス'
        master['B100'] = .1
        master['A101'] = '重複サービス'
        master['B101'] = .2
        book['販売記録']['C6'] = '重複サービス'
        values = self.recalculate(book)
        for row in (5, 6):
            self.assertEqual(values['販売記録'][f'F{row}'].value, '要確認')
            self.assertEqual(values['販売記録'][f'I{row}'].value, '要確認')

    def test_cached_delivery_preserves_formulas_and_validation(self):
        from recalculate import recalculate
        book = build_workbook(rows=20, year=2026)
        expected_formula = book['販売記録']['F5'].value
        expected_validations = book['販売記録'].data_validations.to_tree()
        expected_tables = list(book['販売記録'].tables)
        from xml.etree.ElementTree import tostring
        with tempfile.TemporaryDirectory(prefix='furima-cache-test-') as tmp:
            path = Path(tmp) / 'kit.xlsx'
            book.save(path)
            recalculate(path)
            formulas = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            self.assertEqual(formulas['販売記録']['F5'].value, expected_formula)
            self.assertEqual(tostring(formulas['販売記録'].data_validations.to_tree()),
                             tostring(expected_validations))
            self.assertEqual(list(formulas['販売記録'].tables), expected_tables)
            self.assertEqual(values['販売記録']['I5'].value, 1700)
            self.assertEqual(values['販売記録']['I6'].value, 590)
            self.assertEqual(values['月次集計']['B18'].value, 6500)
            self.assertEqual(values['月次集計']['G18'].value, 2290)
            self.assertEqual(values['採算シミュレータ']['B12'].value, 900)

    def test_simulator_full_fee_has_no_division_error(self):
        book = build_workbook(rows=20, year=2026)
        book['採算シミュレータ']['B9'] = 1
        values = self.recalculate(book)
        self.assertEqual(values['採算シミュレータ']['B12'].value, -1800)
        self.assertNotIsInstance(values['採算シミュレータ']['B14'].value, (int, float))


if __name__ == '__main__':
    unittest.main()
