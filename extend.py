"""既存データを保持し、販売記録・在庫一覧の数式容量を拡張する。"""
from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import MultiCellRange


def extend_workbook(wb, rows):
    """Extend a generated workbook in place without replacing existing cells.

    The last prepared row is the template for newly allocated rows. Source
    range endpoints in formulas, filters, validations, and formatting are
    expanded; existing data and custom cell styles are retained.
    """
    if not isinstance(rows, int) or not 3 <= rows <= 50000:
        raise ValueError('rows は 3～50000 の整数で指定してください。')
    sales = wb['販売記録']
    old_end = sales.max_row
    new_end = rows + 4
    if new_end <= old_end:
        raise ValueError(f'現在の容量 {old_end - 4} 件より大きな --rows を指定してください。')
    inv = wb['在庫一覧']
    if inv.max_row != old_end:
        raise ValueError('販売記録と在庫一覧の容量が一致しません。元のテンプレート構成を確認してください。')
    for ws, formula_cols in [(sales, (6, 9, 12)), (inv, (1, 2, 3, 4, 5, 6))]:
        for r in range(old_end + 1, new_end + 1):
            ws.row_dimensions[r] = copy(ws.row_dimensions[old_end])
            ws.row_dimensions[r].index = r
            for c in range(1, ws.max_column + 1):
                source = ws.cell(old_end, c)
                target = ws.cell(r, c)
                target._style = copy(source._style)
                target.alignment = copy(source.alignment)
                target.protection = copy(source.protection)
                if c in formula_cols and source.data_type == 'f':
                    target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
        ws.auto_filter.ref = re.sub(rf'(?<=[A-Z]){old_end}(?!\d)', str(new_end), ws.auto_filter.ref)
        for validation in ws.data_validations.dataValidation:
            validation.sqref = MultiCellRange(re.sub(rf'(?<=[A-Z]){old_end}(?!\d)', str(new_end), str(validation.sqref)))
        existing = list(ws.conditional_formatting._cf_rules.items())
        ws.conditional_formatting._cf_rules.clear()
        for formatting, rules in existing:
            formatting.sqref = MultiCellRange(re.sub(rf'(?<=[A-Z]){old_end}(?!\d)', str(new_end), str(formatting.sqref)))
            for rule in rules:
                ws.conditional_formatting.add(formatting, rule)
    # Change only the endpoint of explicitly sheet-qualified source ranges.
    pattern = re.compile(r"((?:'販売記録'|販売記録|'在庫一覧'|在庫一覧)!\$[A-Z]+\$\d+:\$[A-Z]+\$)" + str(old_end) + r'(?!\d)')
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    cell.value = pattern.sub(lambda match: match.group(1) + str(new_end), cell.value)
    intro = wb['はじめに']
    for row in intro:
        for cell in row:
            if isinstance(cell.value, str) and '件分の数式を用意' in cell.value:
                cell.value = f'販売記録・在庫一覧は各 {rows:,} 件分の数式を用意。空いている入力行に追記できます。'
    sales['A2'] = f'青＝入力 ｜ 緑＝自動計算 ｜ {rows:,}件分 ｜ 費用ゼロは 0 を入力 ｜ 販売日と仕入日は別に記録'
    if wb.calculation:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--rows', type=int, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error('元ファイル保護のため --output は入力ファイルと異なるパスを指定してください。')
    wb = load_workbook(args.input)
    extend_workbook(wb, args.rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
